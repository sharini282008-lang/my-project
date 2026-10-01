import requests
import streamlit as st
from docx import Document
from docx.shared import Inches, Pt
from reportlab.lib.pagesizes import A4
from reportlab.lib.units import mm
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer
from reportlab.lib.styles import getSampleStyleSheet

BACKEND_URL = "http://127.0.0.1:8000"

st.set_page_config(page_title="LegalEase", page_icon="⚖️", layout="wide")

st.markdown("""
<style>
.main-title { font-size: 42px; font-weight: bold; margin-bottom: 0px; }
.subtitle { font-size: 18px; color: #666666; margin-bottom: 25px; }
.notice { padding: 15px; border-radius: 10px; background-color: #fff3cd;
          border-left: 5px solid #ff9800; margin-bottom: 20px; }
.preview-box { padding: 25px; border-radius: 12px; border: 1px solid #dddddd;
               background-color: #fafafa; line-height: 1.7; }
</style>
""", unsafe_allow_html=True)

st.markdown('<div class="main-title">⚖️ LegalEase</div>', unsafe_allow_html=True)
st.markdown('<div class="subtitle">AI-Powered Legal Document Draft Generator</div>', unsafe_allow_html=True)
st.markdown("""
<div class="notice">
<b>Important:</b> LegalEase generates AI-assisted legal document drafts.
This is not legal advice. Please review the document with a qualified legal professional
before signing or relying on it.
</div>
""", unsafe_allow_html=True)

st.sidebar.title("📄 Document Generator")

document_type = st.sidebar.selectbox(
    "Document Type",
    [
        "Employment Contract",
        "Non-Disclosure Agreement (NDA)",
        "Lease Agreement",
        "Freelance Work Contract",
        "Service Agreement",
        "Offer Letter",
        "General Agreement",
        "Custom"
    ]
)

if document_type == "Custom":
    document_type = st.sidebar.text_input("Custom Document Type", "Legal Agreement")

parties = st.sidebar.text_area(
    "Parties",
    placeholder="Example:\nAlice - Service Provider\nTechNova Inc. - Client",
    height=130
)

terms = st.sidebar.text_area(
    "Terms & Conditions",
    placeholder="Example:\nPayment within 30 days;\nConfidentiality for 2 years;\nEither party may terminate with 15 days notice",
    height=180
)

effective_date = st.sidebar.text_input("Effective Date", "2026-09-25")
jurisdiction = st.sidebar.text_input("Jurisdiction", placeholder="Example: Tamil Nadu, India")
additional_instructions = st.sidebar.text_area(
    "Additional Instructions",
    placeholder="Example: Include signature section.",
    height=100
)

generate_button = st.sidebar.button("✨ Generate Document", type="primary", use_container_width=True)

if generate_button:
    if not parties.strip():
        st.error("Please enter the parties.")
    elif not terms.strip():
        st.error("Please enter the terms and conditions.")
    elif not effective_date.strip():
        st.error("Please enter the effective date.")
    else:
        payload = {
            "document_type": document_type,
            "parties": parties,
            "terms": terms,
            "effective_date": effective_date,
            "jurisdiction": jurisdiction,
            "additional_instructions": additional_instructions
        }
        try:
            with st.spinner("Generating document..."):
                response = requests.post(
                    f"{BACKEND_URL}/api/generate",
                    json=payload,
                    timeout=120
                )
            if response.status_code == 200:
                data = response.json()
                st.session_state["document"] = data["content"]
                st.session_state["document_type"] = data["document_type"]
                st.session_state["model"] = data["model"]
                st.success("✅ Document generated successfully!")
            else:
                try:
                    error_message = response.json().get("detail", response.text)
                except Exception:
                    error_message = response.text
                st.error(f"Backend error: {error_message}")
        except requests.exceptions.ConnectionError:
            st.error("❌ Cannot connect to LegalEase backend.")
            st.info("Make sure the FastAPI backend is running.")
            st.code(".venv\\Scripts\\python.exe -m uvicorn backend.main:app --reload")
        except Exception as error:
            st.error(f"Unexpected error: {error}")

if "document" not in st.session_state:
    st.info("Enter the document information in the left sidebar and click Generate Document.")
else:
    st.subheader("📄 Generated Document")
    edit_mode = st.checkbox("✏️ Edit Document")

    if edit_mode:
        edited_document = st.text_area(
            "Edit your document",
            value=st.session_state["document"],
            height=600
        )
        st.session_state["document"] = edited_document

    document = st.session_state["document"]

    tab1, tab2, tab3 = st.tabs(["📄 Preview", "📝 Raw Text", "⬇️ Downloads"])

    with tab1:
        st.markdown('<div class="preview-box">', unsafe_allow_html=True)
        st.text(document)
        st.markdown('</div>', unsafe_allow_html=True)

    with tab2:
        st.code(document, language="text")

    def create_txt(text):
        return text.encode("utf-8")

    def create_docx(text, document_name):
        doc = Document()
        section = doc.sections[0]
        section.top_margin = Inches(0.75)
        section.bottom_margin = Inches(0.75)
        section.left_margin = Inches(0.85)
        section.right_margin = Inches(0.85)

        normal_style = doc.styles["Normal"]
        normal_style.font.name = "Times New Roman"
        normal_style.font.size = Pt(11)

        title = doc.add_paragraph()
        title.alignment = 1
        run = title.add_run(document_name.upper())
        run.bold = True
        run.font.size = Pt(16)

        for block in text.split("\n\n"):
            block = block.strip()
            if not block:
                continue
            lines = block.splitlines()
            if len(lines) > 1:
                heading = doc.add_paragraph()
                heading_run = heading.add_run(lines[0])
                heading_run.bold = True
                for line in lines[1:]:
                    line = line.strip()
                    if line:
                        paragraph = doc.add_paragraph()
                        paragraph.add_run(line)
            else:
                paragraph = doc.add_paragraph()
                paragraph.add_run(lines[0])

        import io
        buffer = io.BytesIO()
        doc.save(buffer)
        return buffer.getvalue()

    def create_pdf(text, document_name):
        import io
        buffer = io.BytesIO()
        pdf = SimpleDocTemplate(
            buffer,
            pagesize=A4,
            rightMargin=18 * mm,
            leftMargin=18 * mm,
            topMargin=20 * mm,
            bottomMargin=20 * mm
        )
        styles = getSampleStyleSheet()
        story = [Paragraph(document_name.upper(), styles["Title"]), Spacer(1, 12)]

        for block in text.split("\n\n"):
            block = block.strip()
            if not block:
                continue
            for line in block.splitlines():
                line = line.strip()
                if not line:
                    continue
                safe_line = (line.replace("&", "&amp;")
                                  .replace("<", "&lt;")
                                  .replace(">", "&gt;"))
                story.append(Paragraph(safe_line, styles["BodyText"]))
                story.append(Spacer(1, 5))

        pdf.build(story)
        return buffer.getvalue()

    with tab3:
        document_name = st.session_state.get("document_type", "LegalEase Document")
        st.download_button(
            "⬇️ Download TXT",
            data=create_txt(document),
            file_name="legalease_document.txt",
            mime="text/plain",
            use_container_width=True
        )
        st.download_button(
            "⬇️ Download DOCX",
            data=create_docx(document, document_name),
            file_name="legalease_document.docx",
            mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            use_container_width=True
        )
        st.download_button(
            "⬇️ Download PDF",
            data=create_pdf(document, document_name),
            file_name="legalease_document.pdf",
            mime="application/pdf",
            use_container_width=True
        )

st.caption("LegalEase | AI-generated draft | Legal review recommended")
